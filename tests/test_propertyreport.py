"""Property reports (M11, D116, D117): the template, what it prints, the build, checks and prefill. Offline.

The committed template is the only Dynamic file read; the people, property and
pictures here are made up.
"""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from engine.proposal.lightstone import LightstoneFacts
from engine.proposal.store import ProposalStore
from engine.propertyreport import checks, docx_build
from engine.propertyreport.model import (
    Inspection,
    PropertyReport,
    condition_lines,
    extent_text,
    features,
    improvements,
    levies_line,
    owner_line,
    rates_line,
    security_text,
    terms_lines,
    valuation_text,
)
from engine.propertyreport.prefill import apply_facts
from engine.propertyreport.store import ReportStore


def _report(**overrides) -> PropertyReport:
    base = dict(
        dp="9401",
        property_type="house",
        owner_name="Testco (Pty) Ltd",
        owner_id="2014/203299/07",
        legal_description='ERF 99, TOWN "TESTVILLE EXT 1", GAUTENG',
        known_as="1 TEST STREET, TESTVILLE EXT 1, GAUTENG",
        extent="525 m2",
        zoning="RESIDENTIAL",
        local_authority="CITY OF TESTVILLE",
        municipal_valuation=Decimal("850000"),
        title_deed="T12345/2001",
        gps="-25.7,28.2",
        description="The house is a 3-bedroom home in Testville.",
        prepared_by="Test Person",
    )
    base.update(overrides)
    return PropertyReport(**base)


def _picture(path: Path, size=(633, 476)) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (120, 140, 160)).save(path)
    return path.name


def _paragraphs(path: Path):
    doc = Document(str(path))
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


# --- the template ----------------------------------------------------------------------

def test_template_slots():
    text = "\n".join(_paragraphs(docx_build.TEMPLATE))
    assert set(re.findall(r"\{\{[a-z_:]+\}\}", text)) == {
        "{{legal_description}}", "{{known_as}}", "{{owner_line}}", "{{image:cover}}", "{{prepared_by}}",
        "{{prepared_cell}}", "{{prepared_email}}", "{{inspection_date}}", "{{extent}}", "{{zoning}}",
        "{{local_authority}}", "{{municipal_valuation}}", "{{title_deed}}", "{{gps}}", "{{property_heading}}",
        "{{description}}", "{{improvement}}", "{{feature}}", "{{security}}", "{{condition}}", "{{rates_line}}",
        "{{levies_line}}", "{{term}}", "{{viewing_heading}}", "{{viewing_contact}}", "{{conclusion}}",
        "{{signed_by}}", "{{image:photos}}",
    }


def test_template_has_nothing_of_the_master_property_and_the_typos_are_gone():
    text = "\n".join(_paragraphs(docx_build.TEMPLATE)).upper()
    for leftover in ("KADER", "KYALAMI", "TOPHAM", "ST35834", "VENTER", "DECRIPTION", "AREAR"):
        assert leftover not in text, leftover
    assert "PROPERTY DESCRIPTION" in text


def test_template_header_carries_the_properties_letterhead():
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    header = Document(str(docx_build.TEMPLATE)).sections[0].header
    wide = [d for d in header.part.element.iter(qn("w:drawing")) if int(d[0].find(wp + "extent").get("cx")) >= 150 * 36000]
    assert len(wide) == 1
    blob = header.part.related_parts[next(wide[0].iter(a + "blip")).get(embed)].blob
    with Image.open(io.BytesIO(blob)) as im:
        assert im.size == (794, 194)


# --- what it prints ----------------------------------------------------------------------

def test_cover_and_table_values():
    assert owner_line(_report()) == "TESTCO (PTY) LTD, REGISTRATION NUMBER: 2014/203299/07"
    assert owner_line(_report(owner_name="Jane Test", owner_id="7601017486085")) == "JANE TEST ID NUMBER: 7601017486085"
    assert extent_text(_report()) == "525 m²"
    assert valuation_text(_report(municipal_valuation_year=2023)) == "R 850 000.00 (2023 valuation roll)"
    assert valuation_text(_report(municipal_valuation=None)) == "TBC"


def test_the_checklist_prints_only_what_is_filled_in():
    empty = _report()
    assert improvements(empty) == [] and features(empty) == [] and condition_lines(empty) == []
    assert security_text(empty) == ""

    filled = _report(inspection=Inspection(
        bedrooms=3, bathrooms=2, ensuite=1, separate_toilets=1, rooms=["study", "kitchen"], floors="tiled",
        other_improvements=["Walk-in closet", " "], garages=2, patio="covered patio", features=["pool"],
        roof="pitched and tiled", homes=1, farm=["fencing"], security="Remote access", security_items=["alarm", "auto_gates"],
        occupancy="occupied", occupant="tenants", area="medium", electricity="off", meters="prepaid", defects="Broken window",
    ))
    assert improvements(filled) == [
        "3 Bedrooms", "2 Bathrooms (1 en-suite)", "1 Separate toilet", "Kitchen", "Study", "Floors: tiled", "Walk-in closet",
    ]
    assert features(filled) == [
        "2 Garages", "Covered patio", "Swimming pool", "1 Home on the property", "Fencing", "Roof: pitched and tiled",
    ]
    assert security_text(filled) == "Remote access. Alarm system and automated gates."
    assert condition_lines(filled) == [
        "The property is occupied by tenants.", "Area: Medium.", "Services: electricity off, prepaid.", "Defects noted: Broken window.",
    ]


def test_rates_levies_and_terms():
    assert rates_line(_report()) == "We did request a statement and will provide you with same once received."
    assert rates_line(_report(rates_outstanding=Decimal("3610.36"), rates_as_at=date(2026, 7, 6))) == "± R 3 610.36 as at 06/07/2026"
    assert levies_line(_report(levies_outstanding=Decimal("10731.85"), managing_agent="Test Agents")) == "± R 10 731.85\nManaging agent: Test Agents"
    lines = terms_lines(_report(commission_pct=Decimal("7.5"), confirmation_days=30))
    assert lines[1] == "7,5% Commission on the purchase price, plus VAT on the commission payable by the SELLER."
    assert "30 Days confirmation period by the SELLER." in lines
    assert lines[-1] == "Arrear rates and levies for the account of the SELLER."
    assert not any("–" in line or "—" in line for line in lines)


# --- the build ----------------------------------------------------------------------------------

def test_a_full_sectional_report(tmp_path):
    root = tmp_path / "report"
    cover = "cover/" + _picture(root / "cover" / "cover.png", (1138, 900))
    photos = ["photos/" + _picture(root / "photos" / f"p{n}.jpg", (476, 633) if n == 2 else (633, 476)) for n in range(8)]
    report = _report(
        property_type="unit", title_deed="ST65783/2017", inspection_date=date(2026, 9, 10), cover_image=cover, photos=photos,
        inspection=Inspection(bedrooms=2, rooms=["kitchen"], garages=1, security="Remote access", occupancy="vacant"),
        levies_outstanding=Decimal("1200"), managing_agent="Test Agents",
    )
    assert checks.blocking(checks.run(report, root)) == []
    out = docx_build.build(report, root, tmp_path / "9401 - PROPERTY REPORT.docx")
    paras = _paragraphs(out)
    text = "\n".join(paras)

    assert "{{" not in text
    assert "TESTCO (PTY) LTD, REGISTRATION NUMBER: 2014/203299/07" in paras
    assert "Inspection Date: 10 SEPTEMBER 2026" in paras
    assert "Sectional Title Unit:" in paras and "2 Bedrooms" in paras and "1 Garage" in paras
    assert "OCCUPATION AND CONDITION" in paras and "The property is vacant." in paras
    assert any(p.startswith("OUTSTANDING LEVIES") for p in paras) and "Managing agent: Test Agents" in text
    assert any(p.startswith("VIEWING BY APPOINTMENT TO THE UNIT:") for p in paras)
    assert paras[-1] == "" or "Dynamic Auctioneers" in paras

    doc = Document(str(out))
    assert len(list(doc.element.body.iter(qn("w:drawing")))) == 1 + 8
    tables = doc.tables
    assert len(tables) == 1 + 2  # the detail table, then photos six to a page
    assert [len(list(t._tbl.iter(qn("w:drawing")))) for t in tables[1:]] == [6, 2]
    assert tables[0].rows[5].cells[1].text == "R 850 000.00"


def test_a_minimal_freehold_report_leaves_empty_sections_out(tmp_path):
    out = docx_build.build(_report(description=""), tmp_path, tmp_path / "p.docx")
    paras = [p.strip() for p in _paragraphs(out)]
    for heading in ("IMPROVEMENTS:", "EXTERNAL FEATURES:", "SECURITY", "OCCUPATION AND CONDITION", "OUTSTANDING LEVIES:"):
        assert heading not in paras, heading
    assert not any(p.startswith("Inspection Date") for p in paras)
    assert any(p.startswith("VIEWING BY APPOINTMENT:") for p in paras)
    assert "Residential Dwelling:" in paras and "TERMS AND CONDITIONS:" in paras
    assert "{{" not in "\n".join(paras)
    assert len(Document(str(out)).tables) == 1


# --- checks --------------------------------------------------------------------------------------

def test_missing_facts_block():
    fields = {i.field for i in checks.blocking(checks.run(_report(property_type="", owner_name="", title_deed="", prepared_by="")))}
    assert {"property_type", "owner_name", "title_deed", "prepared_by"} <= fields


def test_the_mistakes_in_the_teams_reports_are_warned_about():
    fields = lambda report: [i.field for i in checks.run(report)]  # noqa: E731
    assert "property_type" in fields(_report(title_deed="ST65783/2017"))  # a unit written up as a house (3065)
    assert "property_type" in fields(_report(property_type="vacant", inspection=Inspection(bedrooms=2)))  # 3000
    last_sale = checks.run(_report(municipal_valuation=Decimal("3000000"), lightstone_last_sale=Decimal("3000000"),
                                   lightstone_municipal=Decimal("3002000")))
    assert any("last sale price" in i.message for i in last_sale)  # 3076


def test_a_picture_gone_from_disk_blocks(tmp_path):
    report = _report(photos=["photos/gone.jpg"])
    assert any(i.field == "photos" for i in checks.blocking(checks.run(report, tmp_path)))


# --- prefill and store ------------------------------------------------------------------------------

def test_lightstone_fills_blank_fields_only():
    report = PropertyReport(dp="9402", owner_name="TYPED OWNER")
    note = apply_facts(report, [LightstoneFacts(
        owner_name="TESTCO (PTY) LTD", owner_id="2014/203299/07", title_type="sectional",
        legal_description="SECTION 17 OF THE SCHEME TESTPARK", street_address="UNIT 17 TESTPARK", title_deed="ST65783/2017",
        extent="63 m2", local_authority="CITY OF TESTVILLE", usage="RESIDENTIAL", municipal_valuation="R 850 000",
        municipal_valuation_year="2023", coordinates="-25.7,28.2", last_sale_price="R900 000",
    )])
    assert report.owner_name == "TYPED OWNER" and report.owner_id == "2014/203299/07"
    assert report.property_type == "unit" and report.zoning == "RESIDENTIAL" and report.gps == "-25.7,28.2"
    assert report.municipal_valuation == Decimal("850000") and report.municipal_valuation_year == 2023
    assert report.lightstone_last_sale == Decimal("900000")
    assert note.startswith("Filled from the Lightstone report")


def test_reports_keep_their_own_records(tmp_path):
    db = tmp_path / "engine.db"
    with ReportStore(db) as store:
        store.save(_report(), user="properties@dynamicauctioneers.co.za")
        assert store.get("9401") == _report()
        assert store.list()[0]["seller_name"] == "Testco (Pty) Ltd"
    with ProposalStore(db) as proposals:
        assert proposals.get("9401") is None
