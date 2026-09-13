"""Auction proposals (M9, D103): formatting, OTP terms, checks and the Word build.

Offline and key-free, and no client documents: the OTP used here is written by
the test with the wording every Dynamic OTP shares, and the pictures are plain
generated PNGs. The committed templates are the only real Dynamic files read.
"""

from __future__ import annotations

import re
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from engine.proposal import checks, docx_build, otp_docx
from engine.proposal.model import (
    BudgetLine,
    Erf,
    Proposal,
    Terms,
    budget_rows,
    budget_totals,
    normalise_dp,
    parse_cost,
    pct_text,
    pct_words,
    rands,
    tokens,
)
from engine.proposal.store import ProposalStore

TODAY = date(2026, 6, 1)

OTP_LINES = [
    "INDEX TO DEED OF SALE",
    "AGREEMENT AND CONDITIONS OF SALE",
    "THE JOINT LIQUIDATORS OF: TESTCO (PTY) LTD REGISTRATION NUMBER 2014/203299/07",
    "ERF 99, TOWN “TESTVILLE EXT 1”, GAUTENG",
    "HELD BY TITLE DEED: T12345/2001",
    "ERF 100, TOWN “TESTVILLE EXT 1”, GAUTENG",
    "HELD BY TITLE DEED: T12346/2001",
    "A cash deposit of 10% (Ten Percent) of the PURCHASE PRICE is payable on SIGNATURE DATE.",
    "The commission calculated at 7,5% (SEVEN AND HALF PERCENT) of the purchase price plus VAT, will be payable by the PURCHASER.",
    "This agreement is subject to the Acceptance by the TRUSTEE/SELLER, within a period of 30 (THIRTY) days (the CONFIRMATION PERIOD).",
]


def _write_otp(folder: Path, lines=OTP_LINES) -> Path:
    doc = Document()
    for n, line in enumerate(lines):
        # Numbered like the real contract, so the merge has list numbering to carry.
        doc.add_paragraph(line, style="List Number" if n > 1 else None)
    path = folder / "otp.docx"
    doc.save(str(path))
    return path


def _png(path: Path, size=(1200, 900), colour=(176, 141, 74)) -> str:
    Image.new("RGB", size, colour).save(path)
    return path.name


def _complete(folder: Path) -> Proposal:
    otp = _write_otp(folder)
    terms, notes = otp_docx.read_terms(otp)
    proposal = Proposal(
        dp="9001",
        seller_capacity="liquidation",
        seller_name="TESTCO (PTY) LTD",
        seller_id="2014/203299/07",
        erven=[
            Erf(legal_description="ERF 99, TOWN “TESTVILLE EXT 1”, GAUTENG",
                known_as="1 TEST STREET, TESTVILLE, GAUTENG", title_deed="T12345/2001"),
            Erf(legal_description="ERF 100, TOWN “TESTVILLE EXT 1”, GAUTENG",
                known_as="3 TEST STREET, TESTVILLE, GAUTENG", title_deed="T12346/2001"),
        ],
        auction_date=date(2026, 7, 23),
        time_from="10:00",
        time_to="12:00",
        channel="online",
        administrator="Andrea Oberholzer",
        deadline=date(2026, 6, 11),
        deeds_images=[_png(folder / "deeds-1.png"), _png(folder / "deeds-2.png", (900, 1400))],
        advert_image=_png(folder / "advert.png", (1080, 1350)),
        otp_file=otp.name,
        terms=terms.model_copy(),
        otp_terms=terms.model_copy(),
        otp_notes=notes,
    )
    for line, cost in zip(proposal.budget, ["2500", "302", "2600", "1350", "1943.65", "N/A", "FREE", "1800", "FREE", "N/A", "FREE", "", "FREE"]):
        line.cost = cost
    return proposal


def _doc_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(
        "".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))
    )


# --- formatting --------------------------------------------------------------

def test_money_and_percentages_print_the_way_the_team_writes_them():
    assert rands(Decimal("10495.65")) == "R 10 495.65"
    assert rands(Decimal("1234567")) == "R 1 234 567.00"
    assert rands(Decimal("0")) == "R 0.00"
    assert pct_text(Decimal("7.5")) == "7,5"
    assert pct_text(Decimal("10")) == "10"
    assert pct_words(Decimal("7.5")) == "SEVEN AND A HALF"
    assert pct_words(Decimal("45")) == "FORTY FIVE"
    assert pct_words(Decimal("6.25")) == "SIX POINT TWO FIVE"


def test_budget_totals_match_a_real_proposal_and_blank_lines_are_left_out(tmp_path):
    proposal = _complete(tmp_path)
    total, vat, incl = budget_totals(proposal)
    # DP2821's figures: R 10 495.65 + R 1 574.35 VAT = R 12 070.00.
    assert (total, vat, incl) == (Decimal("10495.65"), Decimal("1574.35"), Decimal("12070.00"))
    rows = budget_rows(proposal)
    assert len(rows) == 12  # the local newspaper line had no cost
    assert {"FREE", "N/A"} <= {r["cost"] for r in rows}


def test_parse_cost_accepts_how_people_type_and_refuses_junk():
    assert parse_cost("R 2 500,50") == (Decimal("2500.50"), None)
    assert parse_cost("1 943.65") == (Decimal("1943.65"), None)
    assert parse_cost("free") == (None, "FREE")
    assert parse_cost("  ") == (None, None)
    with pytest.raises(ValueError):
        parse_cost("about 2k")


def test_dp_numbers_normalise_and_sub_lots_survive():
    assert normalise_dp("DP2990.3") == "2990.3"
    assert normalise_dp("2821") == "2821"
    assert normalise_dp("Moremi") is None


def test_cover_lines_carry_no_dashes(tmp_path):
    values = tokens(_complete(tmp_path))
    assert values["auction_date_line"] == "ONLINE 23 JULY 2026 FROM 10:00 TO 12:00"
    assert values["seller_line"] == "TESTCO (PTY) LTD REGISTRATION NUMBER: 2014/203299/07"
    for text in values.values():
        assert "–" not in text and "—" not in text


# --- OTP terms ---------------------------------------------------------------------

def test_terms_are_read_from_the_otp_wording(tmp_path):
    terms, notes = otp_docx.read_terms(_write_otp(tmp_path))
    assert terms.deposit_pct == Decimal("10")
    assert terms.confirmation_days == 30
    assert terms.commission_pct == Decimal("7.5")
    assert terms.commission_payer == "purchaser"
    assert terms.commission_vat is True
    assert notes == []


def test_a_figure_whose_words_disagree_is_noted_not_resolved(tmp_path):
    lines = [l.replace("7,5% (SEVEN AND HALF PERCENT)", "5% (SIX PERCENT)") for l in OTP_LINES]
    terms, notes = otp_docx.read_terms(_write_otp(tmp_path, lines))
    assert terms.commission_pct == Decimal("5")
    assert any("SIX PERCENT" in n for n in notes)


def test_missing_terms_are_named(tmp_path):
    terms, notes = otp_docx.read_terms(_write_otp(tmp_path, OTP_LINES[:4]))
    assert terms.deposit_pct is None and terms.commission_pct is None
    assert len(notes) == 3


# --- checks ------------------------------------------------------------------------

def _otp_text(folder: Path, proposal: Proposal) -> str:
    return otp_docx.read_text(folder / proposal.otp_file)


def test_a_complete_proposal_has_nothing_blocking(tmp_path):
    proposal = _complete(tmp_path)
    issues = checks.run(proposal, tmp_path, _otp_text(tmp_path, proposal), TODAY)
    assert checks.blocking(issues) == []


def test_checks_catch_the_mistakes_found_on_the_shared_drive(tmp_path):
    proposal = _complete(tmp_path)
    proposal.deadline = date(2025, 5, 7)            # 3018: a 2025 deadline on a 2026 auction
    proposal.terms.deposit_pct = Decimal("20")      # 2887: 20% in the proposal, 10% in the OTP
    proposal.erven[1].title_deed = "T99999/2001"    # the cover and the contract disagree
    issues = checks.blocking(checks.run(proposal, tmp_path, _otp_text(tmp_path, proposal), TODAY))
    fields = {i.field for i in issues}
    assert {"deadline", "deposit_pct", "erf_2_title_deed"} <= fields


def test_a_deadline_after_the_auction_and_missing_files_block(tmp_path):
    proposal = _complete(tmp_path)
    proposal.deadline = date(2026, 8, 1)
    proposal.advert_image = ""
    proposal.otp_file = "nope.docx"
    fields = {i.field for i in checks.blocking(checks.run(proposal, tmp_path, None, TODAY))}
    assert {"deadline", "advert_image", "otp_file"} <= fields


def test_an_unreadable_cost_blocks_and_names_the_line(tmp_path):
    proposal = _complete(tmp_path)
    proposal.budget[2].cost = "about 2k"
    issues = checks.blocking(checks.run(proposal, tmp_path, _otp_text(tmp_path, proposal), TODAY))
    assert any(i.field == "budget_3_cost" and "SAIA" in i.message for i in issues)


# --- the Word build ------------------------------------------------------------------

def test_build_fills_every_token_and_keeps_the_otp_word_for_word(tmp_path):
    proposal = _complete(tmp_path)
    out = docx_build.build(proposal, tmp_path, tmp_path / "out" / "9001 - Auction proposal.docx")
    text = _doc_text(out)

    assert "{{" not in text
    for line in OTP_LINES:
        assert line in text, line
    assert "Liquidated Estate of:" in text
    assert "TESTCO (PTY) LTD REGISTRATION NUMBER: 2014/203299/07" in text
    assert "3 TEST STREET, TESTVILLE, GAUTENG" in text and "\tAND" not in text  # tab is XML, not text
    assert "AND" in text.split("Title deed")[1]
    assert "R 12 070.00" in text and "DEADLINE (11 JUNE 2026)" in text
    assert "A deposit equal to 10% of the purchase price" in text
    assert "within a 30-day period" in text
    assert "7,5% (SEVEN AND A HALF PERCENT)" in text and "PURCHASER commission" in text

    doc = Document(str(out))
    body = doc.element.body
    budget = next(t for t in body.iter(qn("w:tbl")) if "MEDIA" in "".join(x.text or "" for x in t.iter(qn("w:t"))))
    assert len(budget.findall(qn("w:tr"))) == 1 + 12 + 3  # header, printed lines, totals


def test_the_contract_starts_on_a_new_page(tmp_path):
    proposal = _complete(tmp_path)
    out = docx_build.build(proposal, tmp_path, tmp_path / "p.docx")
    doc = Document(str(out))
    first = next(p for p in doc.element.body.iter(qn("w:p"))
                 if "".join(t.text or "" for t in p.iter(qn("w:t"))) == "INDEX TO DEED OF SALE")
    assert first.find(qn("w:pPr")).find(qn("w:pageBreakBefore")) is not None


def test_a_token_without_a_value_refuses_to_build(tmp_path, monkeypatch):
    proposal = _complete(tmp_path)
    real = docx_build.tokens

    def missing_deadline(p):
        values = real(p)
        values.pop("deadline")
        return values

    monkeypatch.setattr(docx_build, "tokens", missing_deadline)
    with pytest.raises(docx_build.TemplateError, match="deadline"):
        docx_build.build(proposal, tmp_path, tmp_path / "p.docx")


def test_templates_carry_only_the_expected_slots():
    def slots(path):
        return set(re.findall(r"\{\{[a-z_:]+\}\}", _doc_text(path)))

    assert slots(docx_build.FRONT) == {
        "{{seller_label}}", "{{seller_line}}", "{{erf_legal}}", "{{erf_known_as}}", "{{erf_title_deed}}",
        "{{auction_date_line}}", "{{auction_venue_line}}", "{{administrator}}", "{{image:deeds}}",
        "{{image:advert}}", "{{auction_starting_line}}", "{{line_media}}", "{{line_description}}",
        "{{line_placement}}", "{{line_cost}}", "{{total}}", "{{vat}}", "{{total_incl}}", "{{deadline}}",
    }
    assert slots(docx_build.BACK) == {
        "{{deposit_pct}}", "{{confirmation_days}}", "{{conduct_channel}}", "{{commission_payer}}",
        "{{commission_pct}}", "{{commission_words}}", "{{commission_vat}}",
    }


def test_every_header_carries_the_properties_letterhead_and_nothing_under_it():
    import io

    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    section = Document(str(docx_build.FRONT)).sections[0]
    blobs = []
    for header in (section.header, section.first_page_header):
        part = header.part
        wide = [
            d for d in part.element.iter(qn("w:drawing"))
            if int(d[0].find(wp + "extent").get("cx")) >= 150 * 36000
        ]
        # One picture only: an older letterhead stacked underneath would show
        # whenever Word's stacking order changed (D106).
        assert len(wide) == 1
        blobs.append(part.related_parts[next(wide[0].iter(a + "blip")).get(embed)].blob)
    assert blobs[0] == blobs[1]
    with Image.open(io.BytesIO(blobs[0])) as im:
        assert im.size == (794, 194)  # Letter Head 2026.png, the properties letterhead


# --- store ---------------------------------------------------------------------------

def test_store_round_trips_and_lists(tmp_path):
    proposal = _complete(tmp_path)
    with ProposalStore(tmp_path / "engine.db") as store:
        store.save(proposal, user="andrea@dynamicauctioneers.co.za")
        again = store.get("9001")
        assert again == proposal
        listed = store.list()
        assert listed[0]["dp"] == "9001" and listed[0]["seller_name"] == "TESTCO (PTY) LTD"
        assert store.delete("9001") and store.get("9001") is None
