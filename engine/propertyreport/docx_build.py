"""Fill the property report Word template (M11, D116, D117).

``templates/property-report.docx`` is the team's master with the DP3060 content
cut out (``scripts/build_report_template.py``). This fills it:

* the bullet lists (improvements, external features, terms) and the occupation
  and condition lines repeat their one slot line per entry;
* a section with nothing to say (no features, no security notes, no levies on a
  freehold property, no inspection date) is removed with its heading, so the
  report never prints an empty heading or "TBC" where nothing was checked;
* the Lightstone picture goes on the cover and the photos follow the text six to
  a page, two across, in borderless tables.

The token filling is the proposal's (M9), so every document fills the same way.
"""

from __future__ import annotations

import copy
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree
from PIL import Image

from engine.proposal.docx_build import _RUN_XPATH, TemplateError, _fill_all, _leftover_tokens, _plain
from engine.propertyreport.model import (
    PropertyReport,
    condition_lines,
    features,
    improvements,
    show_levies,
    terms_lines,
    tokens,
)

TEMPLATE = Path(__file__).resolve().parent / "templates" / "property-report.docx"

# A4 with 1 inch side margins leaves 6.27 inches; under the tall letterhead the
# body is about 8.6 inches high, which three rows of photos fill.
_COVER_W_IN, _COVER_H_IN = 6.27, 5.0
_PHOTO_W_IN, _PHOTO_H_IN = 2.9, 2.2
_COLUMN_IN = 3.13
PER_PAGE = 6


def _para(doc, token: str):
    found = next((p for p in doc.element.body.iter(qn("w:p")) if token in _plain(p)), None)
    if found is None:
        raise TemplateError(f"the template has no {token} slot")
    return found


def _is_blank(el) -> bool:
    return el is not None and el.tag == qn("w:p") and not _plain(el).strip() and not list(el.iter(qn("w:drawing")))


def _remove(el, with_gap: bool = False) -> None:
    gap = el.getnext() if with_gap else None
    el.getparent().remove(el)
    if _is_blank(gap):
        gap.getparent().remove(gap)


def _heading_before(slot, prefix: str):
    el = slot.getprevious()
    while el is not None:
        if _plain(el).strip().upper().startswith(prefix):
            return el
        el = el.getprevious()
    raise TemplateError(f"no {prefix} heading above its slot")


def _repeat(doc, token: str, lines: List[str], heading: Optional[str] = None) -> None:
    slot = _para(doc, "{{%s}}" % token)
    if not lines:
        if heading:
            _remove(_heading_before(slot, heading))
        _remove(slot, with_gap=True)
        return
    pristine = copy.deepcopy(slot)
    _fill_all(slot, {token: lines[0]})
    current = slot
    for line in lines[1:]:
        clone = copy.deepcopy(pristine)
        _fill_all(clone, {token: line})
        current.addnext(clone)
        current = clone


def _drop_section(doc, token: str, heading: str) -> None:
    slot = _para(doc, "{{%s}}" % token)
    _remove(_heading_before(slot, heading))
    _remove(slot, with_gap=True)


def _fit(path: Path, max_w: float, max_h: float):
    with Image.open(path) as im:
        w_px, h_px = im.size
    width = max_w
    height = width * h_px / w_px
    if height > max_h:
        height = max_h
        width = height * w_px / h_px
    return Inches(width), Inches(height)


def _clear(p) -> None:
    for r in p.xpath(_RUN_XPATH):
        r.getparent().remove(r)


def _place_cover(doc, path: Optional[Path]) -> None:
    slot = _para(doc, "{{image:cover}}")
    if path is None:
        _remove(slot)
        return
    _clear(slot)
    paragraph = Paragraph(slot, doc._body)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width, height = _fit(path, _COVER_W_IN, _COVER_H_IN)
    paragraph.add_run().add_picture(str(path), width=width, height=height)


def _place_photos(doc, paths: List[Path]) -> None:
    slot = _para(doc, "{{image:photos}}")
    if not paths:
        _remove(slot)
        return
    _clear(slot)  # the empty slot keeps its page break, and starts the first photo page
    spacer = copy.deepcopy(slot)
    anchor = slot
    for start in range(0, len(paths), PER_PAGE):
        chunk = paths[start:start + PER_PAGE]
        if start:
            # Every later page starts the same way, so no page's photos touch the letterhead rule.
            brk = copy.deepcopy(spacer)
            anchor.addnext(brk)
            anchor = brk
        table = doc.add_table(rows=(len(chunk) + 1) // 2, cols=2)
        table.autofit = False
        for n, path in enumerate(chunk):
            cell = table.cell(n // 2, n % 2)
            cell.width = Inches(_COLUMN_IN)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            width, height = _fit(path, _PHOTO_W_IN, _PHOTO_H_IN)
            paragraph.add_run().add_picture(str(path), width=width, height=height)
        anchor.addnext(table._tbl)  # add_table appended it at the end; this moves it
        anchor = table._tbl
    closing = etree.Element(qn("w:p"))  # Word wants a paragraph after a closing table
    anchor.addnext(closing)


def build(report: PropertyReport, files_root: Path, out_path: Path) -> Path:
    """Write the report ``.docx`` to ``out_path``. Callers run ``checks.run`` first."""
    doc = Document(str(TEMPLATE))
    values = tokens(report)

    if not report.inspection_date:
        _remove(_para(doc, "{{inspection_date}}"))
    if not values["description"]:
        _remove(_para(doc, "{{description}}"))
    _repeat(doc, "improvement", improvements(report), heading="IMPROVEMENTS")
    _repeat(doc, "feature", features(report), heading="EXTERNAL FEATURES")
    if not values["security"]:
        _drop_section(doc, "security", "SECURITY")
    _repeat(doc, "condition", condition_lines(report), heading="OCCUPATION AND CONDITION")
    if not show_levies(report):
        _drop_section(doc, "levies_line", "OUTSTANDING LEVIES")
    _repeat(doc, "term", terms_lines(report), heading="TERMS AND CONDITIONS")
    if not values["viewing_contact"]:
        _remove(_para(doc, "{{viewing_contact}}"), with_gap=True)
    if not values["conclusion"]:
        _drop_section(doc, "conclusion", "CONCLUSION")

    _place_cover(doc, files_root / report.cover_image if report.cover_image else None)
    _place_photos(doc, [files_root / p for p in report.photos])
    _fill_all(doc.element.body, values)

    leftover = _leftover_tokens(doc)
    if leftover:
        raise TemplateError("unfilled tokens: " + ", ".join(leftover))

    core = doc.core_properties
    core.title = f"Property report {report.dp}"
    core.author = "Dynamic Auctioneers"
    core.last_modified_by = "Dynamic Auctioneers"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
